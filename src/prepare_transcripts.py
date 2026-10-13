"""Turn the word-timestamped recall transcripts into per-scene text.

The transcripts (Zenodo 10.5281/zenodo.8208709) are one row per spoken word,
numbered NN01-NN17. Born et al. (2023) documents their exact mapping to
Princeton DataSpace and raw OpenNeuro fMRI IDs. The standalone scene-boundary
table used by this module (Sherlock_Recall_Scene_n50_Onsets.csv) uses the
OpenNeuro numbering, but its subject-1 row mistakenly contains movie rather
than recall onsets. The usable rows therefore follow the published mapping:

  * Matching on recall start/end TR lines up most subjects within 1-8 TR.
  * Slicing each transcript by a candidate's scene windows and scoring the
    resulting text against that scene's description settles the rest: the
    correct pairing scores 0.44-0.54, wrong pairings 0.29-0.42.

The earlier duration/semantic audit independently recovered the same mapping
for subjects 2-16. Subject 1 stays excluded only because its onset row is
invalid, not because its identity is unknown. NN05 has no OpenNeuro/fMRIPrep
counterpart because Princeton subject 05 was omitted from that release.
"""
from __future__ import annotations

import argparse
import glob
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.common import append_data_note, load_config, rel, resolve  # noqa: E402
from src.subject_mapping import OPENNEURO_TO_PRINCETON  # noqa: E402

# OpenNeuro recall-onsets subject -> transcript NN/Princeton number. Subject 1
# is published as NN01 but omitted here because its scene-onset row is corrupt.
SUBJECT_TO_NN = {
    subject: princeton
    for subject, princeton in OPENNEURO_TO_PRINCETON.items()
    if subject != 1
}
UNUSABLE_ONSET_SUBJECTS = [1]
UNUSED_TRANSCRIPTS = [1, 5]

RECALL_FILE = "Sherlock_Recall_Scene_n50_Onsets.csv"
OUT_NAME = "recall_transcripts.csv"
TR_COLUMN = r"Princeton.*TR Num"


def load_words(path: Path) -> pd.DataFrame:
    """One row per spoken word, with its recall TR."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True).active
    header = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
    cols = [i for i, h in enumerate(header) if re.search(TR_COLUMN, h, re.I)]
    if not cols:
        raise ValueError(f"{path.name}: no column matching {TR_COLUMN!r}; "
                         f"found {header}")
    col = cols[0]
    words, trs = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        try:
            trs.append(float(row[col]))
        except (TypeError, ValueError):
            continue
        words.append(str(row[0]))
    return pd.DataFrame({"word": words, "tr": trs})


def build(cfg: dict) -> tuple[pd.DataFrame, list[str]]:
    src = resolve(cfg, "annotations_source")
    tdir = src / "transcripts"
    files = {int(re.search(r"NN(\d+)", Path(f).name).group(1)): Path(f)
             for f in sorted(glob.glob(str(tdir / "NN*.xlsx")))}
    if not files:
        raise FileNotFoundError(
            f"no transcripts in {rel(tdir)}. Fetch them from Zenodo "
            "10.5281/zenodo.8208709 and unzip the 'Sherlock Timestamps' folder there.")

    recall = pd.read_csv(src / RECALL_FILE, encoding="utf-8-sig")
    recall.columns = [c.strip() for c in recall.columns]

    rows = []
    for subject, nn in sorted(SUBJECT_TO_NN.items()):
        if nn not in files:
            continue
        words = load_words(files[nn])
        for _, r in recall[recall["Subject"] == subject].iterrows():
            m = (words["tr"] >= float(r["Onset"])) & (words["tr"] <= float(r["Offset"]))
            if not m.any():
                continue
            rows.append({
                "subject": subject,
                "scene": int(r["Scene"]),
                "transcript_id": f"NN{nn:02d}",
                "n_words": int(m.sum()),
                "text": " ".join(words.loc[m, "word"]),
            })
    out = pd.DataFrame(rows)

    notes = [
        f"- Source: Zenodo `10.5281/zenodo.8208709`, 17 word-timestamped files "
        f"`NN01`-`NN17`; sliced by `{RECALL_FILE}` using the "
        f"`Princeton Dataspace` TR alignment.",
        f"- Born et al. (2023; doi:10.1016/j.dib.2023.109490) publishes the "
        f"OpenNeuro-to-Princeton/NN identity mapping. **{len(SUBJECT_TO_NN)} "
        f"of 16** rows in the standalone recall-onsets table are usable:",
        "",
        "| recall subject | transcript |",
        "|---|---|",
        *[f"| {s} | `NN{n:02d}` |" for s, n in sorted(SUBJECT_TO_NN.items())],
        "",
        f"- OpenNeuro subject {UNUSABLE_ONSET_SUBJECTS[0]} is published as `NN01`/"
        f"Princeton 01, but its row in the standalone CSV contains movie rather "
        f"than recall onsets and is excluded here. The earlier duration/semantic "
        f"audit failed for this reason; it was not evidence against the identity.",
        f"- `NN05` has no OpenNeuro counterpart; Princeton subject 05 was omitted "
        f"from the 16-subject release. `NN01` remains usable through the published "
        f"word timestamps and Chen event labels, which do not rely on the bad CSV row.",
        f"- Result: **{len(out)}** subject x scene text blocks, "
        f"{out['n_words'].sum():,} words total.",
    ]
    return out, notes


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Slice recall transcripts by scene.")
    p.add_argument("--config", default=None)
    args = p.parse_args(argv)
    cfg = load_config(args.config)

    print("=" * 72)
    print("Preparing recall transcripts")
    print("=" * 72)
    table, notes = build(cfg)
    out = resolve(cfg, "annotations_source") / OUT_NAME
    table.to_csv(out, index=False)
    append_data_note(cfg, "Recall transcripts", notes)

    print(f"  subjects mapped : {table['subject'].nunique()} of 16")
    print(f"  scene blocks    : {len(table)}")
    print(f"  words           : {table['n_words'].sum():,}")
    print(f"  median words per recalled scene: {table['n_words'].median():.0f}")
    print(f"\n  Wrote {rel(out)} and recorded the mapping in docs/data_notes.md")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
